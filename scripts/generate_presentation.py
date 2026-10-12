"""
Full Generator for Invoice Decision Intelligence Presentation & Speaker Notes
Creates:
1. PPTX: C:\\Users\\ali\\Downloads\\Invoice_Decision_Intelligence_AIS_Presentation.pptx
2. DOCX: C:\\Users\\ali\\Downloads\\Invoice_Decision_Intelligence_Presentation_Notes.docx
3. PDF:  C:\\Users\\ali\\Downloads\\Invoice_Decision_Intelligence_Presentation.pdf
"""

import os
import sys
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import docx
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = r"C:\Users\ali\Downloads"
PPTX_PATH = os.path.join(OUTPUT_DIR, "Invoice_Decision_Intelligence_AIS_Presentation.pptx")
DOCX_PATH = os.path.join(OUTPUT_DIR, "Invoice_Decision_Intelligence_Presentation_Notes.docx")
PDF_PATH = os.path.join(OUTPUT_DIR, "Invoice_Decision_Intelligence_Presentation.pdf")
ARCH_IMG_PATH = os.path.join(OUTPUT_DIR, "Invoice Decision Intelligence Architecture.png")

# --- COLOR PALETTE ---
C_NAVY = RGBColor(10, 37, 64)        # #0A2540 - Primary Dark
C_SAP_BLUE = RGBColor(0, 102, 204)   # #0066CC - SAP Sapphire Blue
C_TEAL = RGBColor(2, 132, 199)       # #0284C7 - Accent Cyan/Teal
C_SLATE = RGBColor(30, 41, 59)       # #1E293B - Deep Text Slate
C_MUTED = RGBColor(100, 116, 139)    # #64748B - Muted Subtitle
C_LIGHT_BG = RGBColor(248, 250, 252) # #F8FAFC - Card Surface
C_BORDER = RGBColor(226, 232, 240)   # #E2E8F0 - Border Grey
C_WHITE = RGBColor(255, 255, 255)    # #FFFFFF - White
C_GREEN = RGBColor(16, 126, 62)      # #107E3E - Positive Green
C_AMBER = RGBColor(217, 119, 6)      # #D97706 - Warning Amber
C_RED = RGBColor(220, 38, 38)        # #DC2626 - Critical Red
C_PILL_BG = RGBColor(234, 243, 254)  # #EAF3FE - Tag Pill Background
C_PILL_TXT = RGBColor(0, 90, 180)    # #005AB4 - Tag Pill Text

def add_header(slide, tag_text, title_text, subtitle_text):
    """Standard header for content slides."""
    # Tag Pill
    tag_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.42), Inches(2.8), Inches(0.28))
    tag_box.fill.solid()
    tag_box.fill.fore_color.rgb = C_PILL_BG
    tag_box.line.color.rgb = C_PILL_BG
    tf = tag_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = tag_text.upper()
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = C_PILL_TXT
    p.alignment = PP_ALIGN.CENTER

    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.5))
    tf = t_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_NAVY

    # Subtitle
    s_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.26), Inches(11.7), Inches(0.35))
    tf = s_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = subtitle_text
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.color.rgb = C_MUTED

def add_footer(slide, current_slide, total_slides=12):
    """Standard footer across all slides."""
    f_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf = f_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"Invoice Decision Intelligence | AIS Digital Team Presentation                                                                                                               Slide {current_slide} of {total_slides}"
    p.font.name = "Arial"
    p.font.size = Pt(8.5)
    p.font.color.rgb = RGBColor(148, 163, 184)

def add_card(slide, left, top, width, height, border_color=C_BORDER, fill_color=C_LIGHT_BG):
    """Add card shape background."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

print("Helper functions ready")
